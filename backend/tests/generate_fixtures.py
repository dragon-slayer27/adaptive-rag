"""Generates deterministic synthetic fixtures for ingestion testing and verification."""

from pathlib import Path
from docx import Document as DocxWriter
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def generate_sample_pdf(pdf_path: Path) -> None:
    """Generate a valid, deterministic 2-page PDF with extractable text and metadata."""
    writer = PdfWriter()

    font_dict = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    resources = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_dict})
    })

    # Page 1
    page1 = writer.add_blank_page(width=612, height=792)
    stream1 = DecodedStreamObject()
    stream1.set_data(
        b"BT /F1 12 Tf 72 700 Td (Section 1: Agentic RAG System Introduction) Tj "
        b"0 -20 Td (Traditional RAG retrieves documents in a single pass without adaptation.) Tj ET"
    )
    page1[NameObject("/Contents")] = stream1
    page1[NameObject("/Resources")] = resources

    # Page 2
    page2 = writer.add_blank_page(width=612, height=792)
    stream2 = DecodedStreamObject()
    stream2.set_data(
        b"BT /F1 12 Tf 72 700 Td (Section 2: Adaptive Planning and Verification) Tj "
        b"0 -20 Td (Multi-hop evaluation verifies atomic claims against retrieved evidence.) Tj ET"
    )
    page2[NameObject("/Contents")] = stream2
    page2[NameObject("/Resources")] = resources

    writer.add_metadata({
        "/Title": "Agentic RAG Paper",
        "/Author": "Lead Researcher",
    })

    with open(pdf_path, "wb") as f:
        writer.write(f)


def create_all_fixtures(base_dir: Path) -> Path:
    """Create all synthetic test files in base_dir/fixtures."""
    fixtures_dir = base_dir / "fixtures"
    fixtures_dir.mkdir(parents=True, exist_ok=True)

    # 1. Plain text (UTF-8)
    (fixtures_dir / "sample.txt").write_text(
        "Agentic RAG Architecture\n\n"
        "Retrieval-Augmented Generation couples dense retrieval with generation.\n"
        "Adaptive agents plan and execute multi-hop retrieval queries.\n\n"
        "Key capabilities include query decomposition and claim verification.\n",
        encoding="utf-8",
    )

    # 2. Plain text (CP1252 with smart quotes)
    (fixtures_dir / "cp1252_sample.txt").write_bytes(
        b"Legacy Note: \x93Adaptive Retrieval Systems\x94.\n"
        b"System parameters: threshold = 0.85.\n"
    )

    # 3. Markdown with hierarchy
    (fixtures_dir / "sample.md").write_text(
        "# Agentic RAG Overview\n\n"
        "This is the introductory section explaining agentic information retrieval.\n\n"
        "## System Architecture\n\n"
        "The architecture contains ingestion, indexing, and adaptive loop controllers.\n\n"
        "### Ingestion Subsystem\n\n"
        "Parsers convert heterogeneous formats into a canonical schema.\n\n"
        "## Verification and Self-Correction\n\n"
        "Claims are extracted and cross-referenced against evidence.\n",
        encoding="utf-8",
    )

    # 4. DOCX with styles
    doc = DocxWriter()
    doc.core_properties.title = "Agentic RAG Design Note"
    doc.core_properties.author = "System Researcher"
    doc.add_heading("Agentic RAG Design", level=1)
    doc.add_paragraph("Traditional RAG pipelines rely on single-pass retrieval mechanisms.")
    doc.add_heading("Adaptive Retrieval Planning", level=2)
    doc.add_paragraph("Adaptive planners decompose complex queries into verifiable sub-goals.")
    doc.add_heading("Fact Verification", level=2)
    doc.add_paragraph("Extracted atomic claims must be supported by retrieved evidence.")
    doc.save(str(fixtures_dir / "sample.docx"))

    # 5. Multi-page PDF
    generate_sample_pdf(fixtures_dir / "sample.pdf")

    # 6. Empty text file
    (fixtures_dir / "empty.txt").write_bytes(b"")

    # 7. Corrupt PDF file
    (fixtures_dir / "corrupt.pdf").write_bytes(b"%PDF-1.4\nNot a real pdf byte stream\x00\xff")

    # 8. Unsupported file
    (fixtures_dir / "unsupported.xyz").write_text("Unsupported file format content", encoding="utf-8")

    return fixtures_dir


if __name__ == "__main__":
    out_dir = create_all_fixtures(Path(__file__).parent)
    print(f"Fixtures created in: {out_dir}")
