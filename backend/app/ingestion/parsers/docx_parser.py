"""DOCX document parser preserving heading structures and sections."""

from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple
from docx import Document as DocxReader
from docx.opc.exceptions import PackageNotFoundError
from app.ingestion.models import Document, DocumentSection, generate_document_id
from app.ingestion.normalizer import normalize_text
from app.ingestion.parsers.base import BaseParser


class DocxParser(BaseParser):
    """Parser for Microsoft Word (.docx) documents using python-docx."""

    @property
    def supported_extensions(self) -> Set[str]:
        return {".docx"}

    def _extract_sections(self, doc: Any) -> Tuple[Optional[str], List[Tuple[Optional[str], str]]]:
        """Extract sections grouped by heading styles."""
        sections: List[Tuple[Optional[str], List[str]]] = []
        current_title: Optional[str] = None
        current_paras: List[str] = []
        top_title: Optional[str] = None

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else ""

            if style_name.startswith("Heading") or style_name.lower() == "title":
                if style_name.lower() == "title" and top_title is None:
                    top_title = text

                if current_paras or current_title is not None:
                    sections.append((current_title, current_paras))

                current_title = text
                current_paras = []
            else:
                current_paras.append(text)

        if current_paras or current_title is not None:
            sections.append((current_title, current_paras))

        result: List[Tuple[Optional[str], str]] = []
        for heading, paras in sections:
            joined = "\n\n".join(paras)
            normalized = normalize_text(joined)
            if normalized or heading:
                result.append((heading, normalized if normalized else (heading or "")))

        return top_title, result

    def parse(self, file_path: Path, relative_source_path: str, content_hash: str) -> Document:
        raw_bytes = file_path.read_bytes()
        if not raw_bytes:
            raise ValueError(f"File is empty: {file_path.name}")

        try:
            doc = DocxReader(str(file_path))
        except (PackageNotFoundError, Exception) as exc:
            raise ValueError(f"Corrupt or unreadable DOCX '{file_path.name}': {exc}") from exc

        top_title, raw_sections = self._extract_sections(doc)
        if not raw_sections:
            raise ValueError(f"DOCX contains no extractable text: {file_path.name}")

        doc_id = generate_document_id(content_hash)
        sections: List[DocumentSection] = []

        for idx, (heading, sec_text) in enumerate(raw_sections):
            sections.append(
                DocumentSection(
                    document_id=doc_id,
                    section_index=idx,
                    page_number=None,
                    section_title=heading,
                    text=sec_text,
                )
            )

        meta: Dict[str, Any] = {
            "byte_size": len(raw_bytes),
            "paragraph_count": len(doc.paragraphs),
            "section_count": len(sections),
            "headings": [s.section_title for s in sections if s.section_title],
        }

        if doc.core_properties:
            if doc.core_properties.title:
                top_title = doc.core_properties.title.strip()
            if doc.core_properties.author:
                meta["author"] = doc.core_properties.author.strip()

        title = top_title or file_path.stem.replace("_", " ").replace("-", " ").strip().title()

        return Document(
            id=doc_id,
            filename=file_path.name,
            title=title,
            file_type="docx",
            source_path=relative_source_path,
            content_hash=content_hash,
            metadata=meta,
            sections=sections,
        )
